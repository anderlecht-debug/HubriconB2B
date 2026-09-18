"""The spreadsheet and the engine agree on every constant and every demo claim."""
from datetime import timedelta

import pytest
from openpyxl import load_workbook

from hubricon_content import facts as F, playbook_template
from hubricon_engine.models import recovery


@pytest.fixture(scope="module")
def wb(tmp_path_factory):
    p = playbook_template.build(tmp_path_factory.mktemp("xlsx") / "t.xlsx")
    return load_workbook(p)


def test_windows_match_engine(wb):
    ws = wb["Windows"]
    got = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(3, 3 + len(playbook_template.WINDOWS))}
    assert got["warehouse_deadline_days"] == recovery.WAREHOUSE_DEADLINE_DAYS
    assert got["refund_deadline_days"] == recovery.REFUND_DEADLINE_DAYS
    assert got["damaged_return_deadline_days"] == recovery.DAMAGED_RETURN_DEADLINE_DAYS
    assert got["no_cost_value_fraction"] == recovery.NO_COST_VALUE_FRACTION_OF_PRICE


def test_reason_codes_match_engine(wb):
    ws = wb["Reason codes"]
    got = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(2, 2 + len(recovery.REASON_CODES))}
    assert got == recovery.REASON_CODES


def test_demo_claims_evaluate_like_the_engine(wb):
    ws = wb["Windows"]
    hdr_p = next(r for r in range(1, 80) if ws.cell(row=r, column=2).value == "P(approve)")
    hdr_w = next(r for r in range(1, 80) if ws.cell(row=r, column=2).value == "Eligible after (days)")
    p = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(hdr_p + 1, hdr_p + 1 + len(recovery.P_APPROVE))}
    assert p == recovery.P_APPROVE
    rows = {ws.cell(row=r, column=1).value: (ws.cell(row=r, column=2).value, ws.cell(row=r, column=3).value)
            for r in range(hdr_w + 1, hdr_w + 1 + len(recovery.P_APPROVE))}
    data = F.load_data()
    rec = recovery.run(data, today=F.TODAY)
    cs = wb["Claims"]
    for i, c in enumerate(rec["claims"], start=2):
        assert cs.cell(row=i, column=1).value == c["claim_type"]
        ev = cs.cell(row=i, column=3).value.date()
        e, dl = rows[c["claim_type"]]
        assert (ev + timedelta(days=dl)).isoformat() == c["deadline"]
        if c["claim_type"] not in ("damaged_return", "reimbursement_reversal"):
            assert (ev + timedelta(days=e)).isoformat() == c["eligible_from"]
        days_left = (ev + timedelta(days=dl) - F.TODAY).days
        status = ("expired" if F.TODAY > ev + timedelta(days=dl) else "not_yet_eligible" if F.TODAY < ev + timedelta(days=e)
                  else "expiring" if days_left <= recovery.EXPIRING_WITHIN_DAYS else "open")
        assert status == c["status"], (c["claim_type"], c["sku"])
