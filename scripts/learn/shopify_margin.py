"""The Shopify Margin: the course's worked example, its spreadsheet, and the check that holds
the spreadsheet to Hubricon's engine.

    cd engine && uv run --with openpyxl python ../scripts/learn/shopify_margin.py             # figures, then verify
    cd engine && uv run --with openpyxl python ../scripts/learn/shopify_margin.py --publish   # ...and ship the file

FIGURES. One invented Shopify product (a stoneware pour-over set: invented price, compare-at,
weight, costs and orders) priced on the two published cards Hubricon's engine carries in
engine/src/hubricon_engine/cold/priors.py: USPS Ground Advantage commercial rates by zone and
Shopify Payments' plan rates. lib/fees.js prices the same cards in the browser and is held to
priors.py by lib/fees.golden.json, so the course, /call and the engine read one card. Written to
data/learn-shopify-margin.json, which scripts/build-pages.mjs bakes into
learn/shopify-margin.html.

TEMPLATE. learn/files/hubricon-shopify-margin.xlsx: the same arithmetic as formulas.

VERIFY. Random orders, parcels and catalogue rows through priors.py and through the
recalculated sheets, figure by figure. --publish ships the file and
scripts/learn/shopify-margin.stamp.json (the hashes of priors.py, the rate card and the file);
scripts/learn/shopify-margin.test.mjs fails when one moves without a re-run.

Needs LibreOffice (soffice) on PATH.
"""

from __future__ import annotations

import hashlib
import json
import math
import random
import sys
import tempfile
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "engine" / "src"))

from hubricon_engine.cold import priors  # noqa: E402

FILE = ROOT / "learn" / "files" / "hubricon-shopify-margin.xlsx"
FIGURES = ROOT / "data" / "learn-shopify-margin.json"
STAMP = ROOT / "scripts" / "learn" / "shopify-margin.stamp.json"
SOURCES = ["engine/src/hubricon_engine/cold/priors.py", "ratecard.json"]
GRAMS_PER_OZ = 28.349523125
MAX_SHAVEABLE_OZ = 3.0          # lib/fees.js MAX_SHAVEABLE_CARRIER_OZ: further past a line is a redesign
CATALOGUE_DISCOUNT_SHARE = 0.5  # lib/fees.js: half the catalogue below its compare-at makes it a price
MIN_DISCOUNT_SHARE = 0.10       # lib/fees.js: a smaller gap is not an anchor

# ═══ The worked example ══════════════════════════════════════════════════════════════════
# Invented. A single product on a Shopify store with free shipping over $50, paid through
# Shopify Payments on the Basic plan, shipped USPS Ground Advantage from one warehouse.
EXAMPLE = {
    "what": "an invented stoneware pour-over set",
    "price": 42.00,
    "compare_at": 56.00,
    "packed_oz": 16.6,          # the packed parcel: product, box and fill, just past a pound
    "landed_cost": 11.80,
    "packing": 1.40,            # box, fill and the minutes to pack
    "plan": "basic",
    "orders": 640,              # a month, one unit an order
    "free_line": 50.00,         # free shipping over this
    "share_over": 0.55,         # of orders at or over the line: invented
    "catalogue_below": 0.62,    # of the store's variants priced under their own compare-at: invented
}


def sha(path: str) -> str:
    return hashlib.sha256((ROOT / path).read_bytes()).hexdigest()


def row_for(oz: float) -> list[float]:
    """The USPS row a parcel bills at: flat under a pound, then the next whole pound."""
    g = priors.CARRIER_GROUND_COMMERCIAL
    key = 0 if oz < 16 else 16 if oz == 16 else int(math.ceil(oz / 16)) * 16
    return g[key]


def billed(oz: float) -> str:
    return "under a pound" if oz < 16 else f"{math.ceil(oz / 16)} lb"


def figures() -> dict:
    x = EXAMPLE
    g = priors.CARRIER_GROUND_COMMERCIAL
    fee = priors.payments_fee(x["plan"], x["price"])
    now, under = row_for(x["packed_oz"]), g[0]
    step = priors.CARRIER_GROUND_USD[16]
    zones = list(range(1, 9))
    contrib = [x["price"] - fee - x["packing"] - x["landed_cost"] - lab for lab in now]
    contrib_under = [x["price"] - fee - x["packing"] - x["landed_cost"] - lab for lab in under]
    roas = [x["price"] / c for c in contrib]
    over = round(x["orders"] * x["share_over"])
    anchor = round(x["compare_at"] - x["price"], 2)
    r2 = lambda v, d=2: round(float(v), d)
    return {
        "about": "The Shopify Margin's worked example: an invented Shopify product priced on USPS Ground Advantage's and Shopify Payments' published rates, as Hubricon's engine carries them (scripts/learn/shopify_margin.py).",
        "example": {k: v for k, v in x.items()},
        "cards": {"usps_source": priors.CARRIER_SOURCE if hasattr(priors, "CARRIER_SOURCE") else None,
                  "payments_source": priors.PAYMENTS_SOURCE, "payments": {k: list(v) for k, v in priors.SHOPIFY_PAYMENTS.items()},
                  "gateway_surcharge": priors.THIRD_PARTY_GATEWAY_SURCHARGE, "usps": {str(k): v for k, v in g.items()}},
        "order": {
            "fee": r2(fee, 4), "label_now": [r2(v) for v in now], "label_under": [r2(v) for v in under], "billed": billed(x["packed_oz"]),
            "contribution": [r2(v, 4) for v in contrib], "contribution_under": [r2(v, 4) for v in contrib_under],
            "contribution_lo": r2(min(contrib), 4), "contribution_hi": r2(max(contrib), 4),
            "fixed_before_label": r2(x["price"] - fee - x["packing"] - x["landed_cost"], 4),
        },
        "pound": {"edge": 16, "over_by": r2(x["packed_oz"] - 16, 1), "step": list(step),
                  "month_lo": r2(step[0] * x["orders"]), "month_hi": r2(step[1] * x["orders"]),
                  "trim_to": 15.9},
        "free": {"orders_over": over, "subsidy_lo": r2(over * min(now)), "subsidy_hi": r2(over * max(now)),
                 "label_lo": r2(min(now)), "label_hi": r2(max(now))},
        "anchor": {"per_unit": anchor, "share": r2(anchor / x["compare_at"], 4), "catalogue_share": x["catalogue_below"],
                   "policy": x["catalogue_below"] >= CATALOGUE_DISCOUNT_SHARE,
                   "shown_month": r2(anchor * x["orders"]), "step5": r2(x["price"] * 1.05)},
        "ads": {"roas_lo": r2(min(roas), 2), "roas_hi": r2(max(roas), 2), "cpa_lo": r2(min(contrib)), "cpa_hi": r2(max(contrib)),
                "roas_under_lo": r2(min(x["price"] / c for c in contrib_under), 2), "roas_under_hi": r2(max(x["price"] / c for c in contrib_under), 2)},
        "zones": zones,
        "sources": {p: sha(p) for p in SOURCES},
    }


# ═══ The spreadsheet ═════════════════════════════════════════════════════════════════════
INPUT = PatternFill("solid", fgColor="FFF2CC")
TITLE = Font(name="Arial", bold=True, size=16)
HEAD = Font(name="Arial", bold=True, size=11)
BODY = Font(name="Arial", size=10)
BOLD = Font(name="Arial", bold=True, size=10)
NOTE = Font(name="Arial", size=9, italic=True, color="555555")
USD, PCT, PCT1, NUM1, NUM2, INT = '"$"#,##0.00', "0%", "0.0%", "0.0", "0.00", "#,##0"
CATALOGUE_ROWS = 200


def put(ws, ref, value, font=BODY, fmt=None, fill=None, bold=False):
    c = ws[ref]
    c.value = value
    c.font = BOLD if bold else font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def cards_sheet(ws):
    """The two published cards, as the engine carries them. Every other sheet looks up here."""
    ws.title = "Cards"
    put(ws, "A1", "The cards: USPS Ground Advantage commercial and Shopify Payments", TITLE)
    put(ws, "A2", "As Hubricon's engine carries them. Check the live rates before you act: USPS reprices about once a year, Shopify by plan.", NOTE)
    put(ws, "A4", "USPS Ground Advantage, commercial: the row a parcel bills at, by zone", HEAD)
    put(ws, "A5", "Row (oz)", BOLD)
    for z in range(8):
        put(ws, f"{chr(66 + z)}5", f"Zone {z + 1}", BOLD)
    for i, (oz, row) in enumerate(sorted(priors.CARRIER_GROUND_COMMERCIAL.items())):
        r = 6 + i
        put(ws, f"A{r}", oz)
        for z, v in enumerate(row):
            put(ws, f"{chr(66 + z)}{r}", v, fmt=USD)
    put(ws, "A10", "Row 0 is anything under a pound; row 16 a parcel of exactly one pound; then every part pound rounds up to the next row.", NOTE)
    put(ws, "A12", "Shopify Payments, online card, United States", HEAD)
    for c, h in zip("ABCD", ["Plan", "Rate", "Fixed a charge", "Surcharge if you use another gateway"]):
        put(ws, f"{c}13", h, BOLD)
    for i, (plan, (rate, fixed)) in enumerate(priors.SHOPIFY_PAYMENTS.items()):
        r = 14 + i
        put(ws, f"A{r}", plan)
        put(ws, f"B{r}", rate, fmt=PCT1)
        put(ws, f"C{r}", fixed, fmt=USD)
        put(ws, f"D{r}", priors.THIRD_PARTY_GATEWAY_SURCHARGE[plan], fmt=PCT1)
    put(ws, "A19", "Sources: " + priors.PAYMENTS_SOURCE + "; USPS Notice 123, effective 2026-07-12 (pe.usps.com).", NOTE)
    ws.column_dimensions["A"].width = 30
    for z in range(8):
        ws.column_dimensions[chr(66 + z)].width = 12


ROW_KEY = "IF({w}<16,0,IF({w}=16,16,CEILING({w}/16,1)*16))"
LABEL = "INDEX(Cards!$B$6:$I$9,MATCH(" + ROW_KEY + ",Cards!$A$6:$A$9,0),{z})"


def order_sheet(ws, x):
    """Lessons 1 and 5: where one order's money goes, by zone, and the most an ad sale can cost."""
    ws.title = "Your order"
    put(ws, "A1", "Your order: where one order's money goes, and the most an ad can cost", TITLE)
    put(ws, "A2", "Yellow cells are yours. One unit an order, free shipping: the label is yours to pay. Shopify Payments charges on the whole charge.", NOTE)
    for r, label, v, fmt in [(4, "Price", x["price"], USD), (5, "Packed weight, ounces (weigh a packed parcel)", x["packed_oz"], NUM1),
                             (6, "Landed cost a unit", x["landed_cost"], USD), (7, "Packing a unit: box, fill, minutes", x["packing"], USD),
                             (8, "Shopify plan (basic, grow, advanced, plus)", x["plan"], None)]:
        put(ws, f"A{r}", label)
        put(ws, f"B{r}", v, fmt=fmt, fill=INPUT)
    put(ws, "A10", "Shopify Payments a charge", BOLD)
    put(ws, "B10", "=ROUND(B4*VLOOKUP(B8,Cards!$A$14:$C$17,2,FALSE)+VLOOKUP(B8,Cards!$A$14:$C$17,3,FALSE),4)", fmt=USD, bold=True)
    put(ws, "A11", "The parcel bills as (oz row)")
    put(ws, "B11", "=" + ROW_KEY.format(w="B5"), fmt=INT)
    put(ws, "A12", "Before the label: price − Shopify Payments − packing − landed cost")
    put(ws, "B12", "=B4-B10-B7-B6", fmt=USD)
    for c, h in zip("ABCD", ["Zone", "The label", "Kept before ads (the most one ad sale can cost)", "Break-even ROAS (price ÷ kept)"]):
        put(ws, f"{c}14", h, BOLD)
    for z in range(1, 9):
        r = 14 + z
        put(ws, f"A{r}", z)
        put(ws, f"B{r}", "=" + LABEL.format(w="$B$5", z=z), fmt=USD)
        put(ws, f"C{r}", f"=$B$12-B{r}", fmt=USD)
        put(ws, f"D{r}", f'=IF(C{r}>0,$B$4/C{r},"loses on every order")', fmt=NUM2)
    put(ws, "A24", "An ad that returns less than the break-even ROAS loses money on every sale it buys, before your fixed costs.", NOTE)
    ws.column_dimensions["A"].width = 56
    for c in "BCD":
        ws.column_dimensions[c].width = 22


def pound_sheet(ws, x):
    """Lesson 2: one parcel against the pound line."""
    ws.title = "Pound line"
    put(ws, "A1", "Pound line: what a parcel just past a pound costs, and what trimming it saves", TITLE)
    put(ws, "A2", "Yellow cells are yours. USPS rounds every part pound up: 16.1 ounces bills as 2 lb.", NOTE)
    for r, label, v, fmt in [(4, "Packed weight, ounces", x["packed_oz"], NUM1), (5, "Orders a month (one parcel each)", x["orders"], INT)]:
        put(ws, f"A{r}", label)
        put(ws, f"B{r}", v, fmt=fmt, fill=INPUT)
    put(ws, "A7", "Bills as (oz row)")
    put(ws, "B7", "=" + ROW_KEY.format(w="B4"), fmt=INT)
    put(ws, "A8", "The line below it (oz)")
    put(ws, "B8", '=IF(B4<=16,"",FLOOR(B4-0.000001,16))', fmt=INT)
    put(ws, "A9", "Ounces past the line")
    put(ws, "B9", '=IF(B8="","",B4-B8)', fmt=NUM1)
    put(ws, "A10", "Worth trimming? (within 3 ounces)")
    put(ws, "B10", f'=IF(B8="","Under a pound: no line below",IF(B9<={MAX_SHAVEABLE_OZ},"Yes","No: that is a redesign"))')
    for c, h in zip("ABCDE", ["Zone", "Now", "Under the line", "Saved a parcel", "Saved a month"]):
        put(ws, f"{c}12", h, BOLD)
    under_row = "IF($B$8=16,0,$B$8)"
    for z in range(1, 9):
        r = 12 + z
        put(ws, f"A{r}", z)
        put(ws, f"B{r}", "=" + LABEL.format(w="$B$4", z=z), fmt=USD)
        put(ws, f"C{r}", f'=IF($B$8="","",INDEX(Cards!$B$6:$I$9,MATCH({under_row},Cards!$A$6:$A$9,0),{z}))', fmt=USD)
        put(ws, f"D{r}", f'=IF($B$8="","",B{r}-C{r})', fmt=USD)
        put(ws, f"E{r}", f'=IF($B$8="","",D{r}*$B$5)', fmt=USD)
    put(ws, "A22", "Your zone mix is in Shopify's shipping reports or your label history; until you have it, the range across zones is the honest figure.", NOTE)
    ws.column_dimensions["A"].width = 40
    for c in "BCDE":
        ws.column_dimensions[c].width = 16


def free_sheet(ws, x):
    """Lesson 3: the free-shipping line."""
    ws.title = "Free shipping"
    put(ws, "A1", "Free shipping: what the line you drew costs a month", TITLE)
    put(ws, "A2", "Yellow cells are yours. Over the line you pay the label; under it the customer does.", NOTE)
    for r, label, v, fmt in [(4, "Orders a month", x["orders"], INT), (5, "Share of orders at or over the line", x["share_over"], PCT),
                             (6, "Average label on an order over the line", "=AVERAGE('Your order'!B15:B22)", USD)]:
        put(ws, f"A{r}", label)
        put(ws, f"B{r}", v, fmt=fmt, fill=INPUT)
    put(ws, "A8", "Orders over the line", BOLD)
    put(ws, "B8", "=ROUND(B4*B5,0)", fmt=INT)
    put(ws, "A9", "What the line costs a month", BOLD)
    put(ws, "B9", "=B8*B6", fmt=USD, bold=True)
    put(ws, "A10", "Extra kept a month the line must bring to pay for itself")
    put(ws, "B10", "=B9", fmt=USD)
    put(ws, "A12", "The line pays only if the orders and the larger baskets it brings keep at least what it costs. Test it the way The Price Curve tests a price: move it, and measure.", NOTE)
    ws.column_dimensions["A"].width = 56
    ws.column_dimensions["B"].width = 18


def catalogue_sheet(ws, rows=None):
    """Lesson 6: every variant, pasted from Shopify's Products export."""
    ws.title = "Catalogue"
    put(ws, "A1", "Catalogue: every variant, from your Products export", TITLE)
    put(ws, "A2", "Paste from Shopify Admin, Products, Export: Variant SKU, Variant Price, Variant Compare At Price, Variant Grams. Grams are converted to ounces; add your box and fill if the grams are the product alone.", NOTE)
    heads = ["Variant SKU", "Price", "Compare-at", "Grams", "Ounces", "Bills as (oz row)", "Ounces past a line",
             "Saved a parcel, nearest zone", "Saved a parcel, farthest zone", "Below its compare-at by", "Share off"]
    for i, h in enumerate(heads):
        put(ws, f"{chr(65 + i)}4", h, BOLD)
    put(ws, "M4", "Share of priced variants below their compare-at", BOLD)
    put(ws, "M5", f'=IFERROR(COUNTIF(K5:K{4 + CATALOGUE_ROWS},">0")/COUNT(B5:B{4 + CATALOGUE_ROWS}),"")', fmt=PCT)
    put(ws, "M6", f'=IF(M5="","",IF(M5>={CATALOGUE_DISCOUNT_SHARE},"Half or more: the sale is your price","Under half: a promotion, not a policy"))')
    for i in range(CATALOGUE_ROWS):
        r = 5 + i
        x = rows[i] if rows and i < len(rows) else None
        for c, k, fmt in [("A", "sku", None), ("B", "price", USD), ("C", "compare_at", USD), ("D", "grams", NUM1)]:
            put(ws, f"{c}{r}", x[k] if x else None, fmt=fmt, fill=INPUT)
        w = f"E{r}"
        put(ws, f"E{r}", f'=IF(D{r}="","",ROUND(D{r}/{GRAMS_PER_OZ},2))', fmt=NUM2)
        put(ws, f"F{r}", f'=IF(E{r}="","",{ROW_KEY.format(w=w)})', fmt=INT)
        put(ws, f"G{r}", f'=IF(OR(E{r}="",E{r}<=16),"",E{r}-FLOOR(E{r}-0.000001,16))', fmt=NUM2)
        near = f'INDEX(Cards!$B$6:$I$9,MATCH(F{r},Cards!$A$6:$A$9,0),{{z}})-INDEX(Cards!$B$6:$I$9,MATCH(IF(F{r}-16=16,0,F{r}-16),Cards!$A$6:$A$9,0),{{z}})'
        ok = f'AND(G{r}<>"",G{r}<={MAX_SHAVEABLE_OZ},F{r}<=48)'
        # the step across zones: the nearest zone's saving is the smallest, the farthest's the largest, as on the card
        put(ws, f"H{r}", f'=IF({ok},MIN({",".join(near.format(z=z) for z in range(1, 9))}),"")', fmt=USD)
        put(ws, f"I{r}", f'=IF({ok},MAX({",".join(near.format(z=z) for z in range(1, 9))}),"")', fmt=USD)
        put(ws, f"J{r}", f'=IF(AND(B{r}<>"",C{r}<>"",C{r}>B{r}),C{r}-B{r},"")', fmt=USD)
        put(ws, f"K{r}", f'=IF(J{r}="","",IF(ROUND(J{r}/C{r},4)>={MIN_DISCOUNT_SHARE},ROUND(J{r}/C{r},4),0))', fmt=PCT1)
    for i, wd in enumerate([18, 10, 12, 10, 10, 14, 14, 16, 16, 16, 10]):
        ws.column_dimensions[chr(65 + i)].width = wd
    ws.column_dimensions["M"].width = 30


def start_sheet(ws):
    ws.title = "Start here"
    lines = [("The Shopify Margin", TITLE), ("A free course from Hubricon: hubricon.com/learn/shopify-margin", BODY), ("", BODY),
             ("What this file does", HEAD),
             ("Your order: where one order's money goes, zone by zone, and the most one ad sale can cost (break-even ROAS).", BODY),
             ("Pound line: what a parcel just past a pound costs, and what trimming it saves a month.", BODY),
             ("Free shipping: what the line you drew costs a month.", BODY),
             ("Catalogue: every variant from your Products export, with the parcels just past a line and the prices under their own compare-at.", BODY),
             ("Cards: USPS Ground Advantage and Shopify Payments, as Hubricon's engine carries them.", BODY), ("", BODY),
             ("How to use it", HEAD), ("Yellow cells are yours; everything else is a formula you can read.", BODY),
             (f"It opens on the course's example: {EXAMPLE['what']}. Its price, weight, costs and orders are invented; the cards are the published ones.", BODY),
             ("", BODY), ("Where the arithmetic comes from", HEAD),
             ("Hubricon's engine (cold/priors.py) and lib/fees.js, the same cards /call reads. Checked against the engine before this file was published.", BODY)]
    for i, (t, f) in enumerate(lines):
        put(ws, f"A{i + 1}", t, f)
    ws.column_dimensions["A"].width = 120


def build(path: Path, catalogue=None) -> Workbook:
    wb = Workbook()
    start_sheet(wb.active)
    order_sheet(wb.create_sheet(), EXAMPLE)
    pound_sheet(wb.create_sheet(), EXAMPLE)
    free_sheet(wb.create_sheet(), EXAMPLE)
    catalogue_sheet(wb.create_sheet(), catalogue)
    cards_sheet(wb.create_sheet())
    wb.properties.creator = "Hubricon"
    wb.properties.title = "The Shopify Margin"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return wb


def recalc(path: Path) -> None:
    sys.path.insert(0, str(ROOT / "scripts" / "learn"))
    from verify_fee_staircase import recalc as lo_recalc
    lo_recalc(path)


# ═══ The check ═══════════════════════════════════════════════════════════════════════════
def verify(fig: dict, cases: int = 30) -> int:
    rng = random.Random(20261003)
    problems: list[str] = []
    orders = [{"price": round(rng.uniform(9, 120), 2), "oz": round(rng.uniform(2, 47.9), 1), "landed": round(rng.uniform(1, 30), 2),
               "packing": round(rng.uniform(0.3, 3), 2), "plan": rng.choice(list(priors.SHOPIFY_PAYMENTS))} for _ in range(cases)]
    orders += [{"price": 30.0, "oz": w, "landed": 8.0, "packing": 1.0, "plan": "basic"} for w in (15.99, 16.0, 16.01, 32.0, 32.5, 48.0)]
    catalogue = []
    for i in range(60):
        price = round(rng.uniform(8, 90), 2)
        ca = rng.choice([None, round(price * rng.uniform(1.02, 1.6), 2), round(price * 0.9, 2)])
        catalogue.append({"sku": f"V{i}", "price": price, "compare_at": ca, "grams": round(rng.uniform(40, 1500), 1) if rng.random() > 0.1 else None})
    with tempfile.TemporaryDirectory(prefix="hubricon-sm-") as tmp:
        path = Path(tmp) / "check.xlsx"
        wb = build(path, catalogue)
        for k, o in enumerate(orders):
            ws = wb.copy_worksheet(wb["Your order"])
            ws.title = f"O{k}"
            for ref, v in (("B4", o["price"]), ("B5", o["oz"]), ("B6", o["landed"]), ("B7", o["packing"]), ("B8", o["plan"])):
                ws[ref].value = v
            pw = wb.copy_worksheet(wb["Pound line"])
            pw.title = f"P{k}"
            pw["B4"].value = o["oz"]
            pw["B5"].value = 100
        wb.save(path)
        recalc(path)
        got = load_workbook(path, data_only=True)
        for k, o in enumerate(orders):
            ws, pw = got[f"O{k}"], got[f"P{k}"]
            fee = priors.payments_fee(o["plan"], o["price"])
            if abs(ws["B10"].value - fee) > 1.5e-4:   # the engine rounds to four places; the sheet shows cents
                problems.append(f"O{k}: fee sheet {ws['B10'].value} engine {fee}")
            row = row_for(o["oz"])
            for z in range(8):
                lab = ws[f"B{15 + z}"].value
                if abs(lab - row[z]) > 1e-9:
                    problems.append(f"O{k} zone {z + 1}: label sheet {lab} engine {row[z]}")
                kept = o["price"] - fee - o["packing"] - o["landed"] - row[z]
                if abs(ws[f"C{15 + z}"].value - kept) > 1.5e-4:
                    problems.append(f"O{k} zone {z + 1}: kept sheet {ws[f'C{15 + z}'].value} expected {kept}")
            # the pound line: the step to the band below, against the engine's own carrier_rows
            if 16 < o["oz"] <= 48:
                edge = int(math.floor((o["oz"] - 1e-6) / 16)) * 16
                rows = priors.carrier_rows(edge)
                if rows:
                    for z in range(8):
                        want = rows[0][z] - rows[1][z]
                        got_v = pw[f"D{13 + z}"].value
                        if abs(got_v - want) > 1e-9:
                            problems.append(f"P{k} ({o['oz']} oz) zone {z + 1}: saved sheet {got_v} engine {want}")
        cat = got["Catalogue"]
        below = [c for c in catalogue if c["compare_at"] and c["compare_at"] > c["price"]]
        share = len([c for c in below if round((c["compare_at"] - c["price"]) / c["compare_at"], 4) >= MIN_DISCOUNT_SHARE]) / len(catalogue)
        if abs(cat["M5"].value - share) > 1e-9:
            problems.append(f"Catalogue share sheet {cat['M5'].value} expected {share}")
        for i, c in enumerate(catalogue):
            r = 5 + i
            if not c["grams"]:
                continue
            oz = round(c["grams"] / GRAMS_PER_OZ, 2)
            if 16 < oz:
                edge = int(math.floor((oz - 1e-6) / 16)) * 16
                over = oz - edge
                steps = priors.CARRIER_GROUND_USD.get(edge)
                want = steps if (steps and over <= MAX_SHAVEABLE_OZ) else None
                h, ii = cat[f"H{r}"].value, cat[f"I{r}"].value
                got_pair = (h, ii) if isinstance(h, (int, float)) else None
                if (want is None) != (got_pair is None) or (want and (abs(got_pair[0] - want[0]) > 1e-6 or abs(got_pair[1] - want[1]) > 1e-6)):
                    problems.append(f"Catalogue {c['sku']} ({oz} oz): saved sheet {got_pair} engine {want}")
        o = got["Your order"]
        for z in range(8):
            if abs(o[f"C{15 + z}"].value - fig["order"]["contribution"][z]) > 1e-3:
                problems.append(f"Example zone {z + 1}: kept sheet {o[f'C{15 + z}'].value} figures {fig['order']['contribution'][z]}")
    for p in problems:
        print("  ✗", p)
    print(f"{2 * len(orders) + len(catalogue)} golden cases and the worked example: {'all agree' if not problems else f'{len(problems)} disagree'}")
    return len(problems)


if __name__ == "__main__":
    fig = figures()
    FIGURES.write_text(json.dumps(fig, indent=1) + "\n")
    o, p, a, ads = fig["order"], fig["pound"], fig["anchor"], fig["ads"]
    print(f"fee {o['fee']}, kept {o['contribution_lo']}–{o['contribution_hi']}, step {p['step']} → {p['month_lo']}–{p['month_hi']}/month, "
          f"free line {fig['free']['subsidy_lo']}–{fig['free']['subsidy_hi']}, anchor {a['per_unit']} ({a['share']}), ROAS {ads['roas_lo']}–{ads['roas_hi']}")
    print(f"wrote {FIGURES.relative_to(ROOT)}")
    if "--figures-only" in sys.argv:
        raise SystemExit(0)
    if verify(fig):
        raise SystemExit(1)
    if "--publish" in sys.argv:
        build(FILE)
        recalc(FILE)
        STAMP.write_text(json.dumps({"file": str(FILE.relative_to(ROOT)), "file_sha256": hashlib.sha256(FILE.read_bytes()).hexdigest(),
                                     "sources_sha256": fig["sources"], "verified_on": date.today().isoformat()}, indent=2) + "\n")
        print(f"published {FILE.relative_to(ROOT)} and {STAMP.relative_to(ROOT)}")
