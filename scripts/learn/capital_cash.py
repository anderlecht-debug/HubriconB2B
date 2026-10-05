"""Capital & Cash: the course's worked example, computed by Hubricon's engine.

    cd engine && uv run python ../scripts/learn/capital_cash.py            # figures -> data/learn-capital-cash.json

FIGURES. The worked example is the same invented garlic press The Price Curve prices: an
invented sales history, invented costs and an invented bank balance, with Amazon's published
fees. Every figure the course prints from it is computed here by the engine itself and written
to data/learn-capital-cash.json, which scripts/build-pages.mjs bakes into
learn/capital-and-cash.html. Nothing on the page is typed.

    the reorder point, order size and stockout odds   models/inventory_sim.py  run, demand_sf, demand_quantile
    the cash path, its low point and the tail          models/cashflow.py       run (10,000 paths)
    the service level each SKU earns                   models/inventory_econ.py critical_fractile, demand_over_cycle
    hold or liquidate the excess                       models/inventory_econ.py hold_vs_liquidate
    the payout cycle                                   channels.py              PAYOUT_CYCLE_DAYS

The second SKU (an invented silicone spatula set) is there to show a thin margin beside a fat
one. The stamp beside the data pins the hashes of the engine modules it was computed with;
scripts/learn/capital-cash.test.mjs fails when one moves without a re-run.
"""

from __future__ import annotations

import hashlib
import json
import math
import sys
from datetime import date
from pathlib import Path

import numpy as np

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "engine" / "src"))

from hubricon_engine import channels  # noqa: E402
from hubricon_engine.cold import priors  # noqa: E402
from hubricon_engine.models import cashflow, inventory_econ, inventory_sim  # noqa: E402
from hubricon_engine.models import fee_schedule as fees  # noqa: E402

FIGURES = ROOT / "data" / "learn-capital-cash.json"
ENGINE_SOURCES = [
    "engine/src/hubricon_engine/models/inventory_sim.py",
    "engine/src/hubricon_engine/models/inventory_econ.py",
    "engine/src/hubricon_engine/models/cashflow.py",
    "engine/src/hubricon_engine/models/markdown.py",
    "engine/src/hubricon_engine/models/fee_schedule.py",
    "engine/src/hubricon_engine/channels.py",
]
TODAY = date(2026, 10, 1)              # the example's "today": the snapshot is the day before
PRICED_ON = date(2026, 9, 30)          # the card the fulfilment fee is read from (as The Price Curve)

# ═══ The worked example ══════════════════════════════════════════════════════════════════
# Invented. Twelve half-months of sales drawn around The Price Curve's 523 units a month, so
# the two courses describe one listing; the seed is one whose draw lands on that rate.
GARLIC = {
    "what": "an invented stainless garlic press",
    "sku": "GP-1",
    "category": "home & kitchen",
    "tier": "large_standard",
    "weight_oz": 12.0,
    "volume_cuft": 0.046,          # a 9 × 3.5 × 2.5 inch box
    "price": 24.99,
    "landed_cost": 6.20,
    "lead_days": 75,               # 30 days to make, 35 to ship, 10 to check in: invented, inside the 66–119 days sourced for China to FBA
    "on_hand": 2500,
    "ad_month": 900.0,
    "cash": 11000.0,
    "fixed_month": 4000.0,
    "history_seed": 20261371,
    "history_rate": 17.2,          # The Price Curve's 523 a month, a day
    "history_cv": 0.18,
}
SPATULA = {
    "what": "an invented silicone spatula set",
    "sku": "SP-3",
    "category": "home & kitchen",
    "tier": "large_standard",
    "weight_oz": 9.0,
    "volume_cuft": 0.031,
    "price": 12.99,
    "landed_cost": 5.40,
    "lead_days": 75,
    "rate": 3.0,                   # units a day
    "rate_sd": 0.6,
    "position": 1400,              # what is on hand: far more than it sells
    "age_days": 200,
}
HALF_MONTHS = [(4, 1, 15), (4, 16, 30), (5, 1, 15), (5, 16, 31), (6, 1, 15), (6, 16, 30),
               (7, 1, 15), (7, 16, 31), (8, 1, 15), (8, 16, 31), (9, 1, 15), (9, 16, 30)]
CASH_PATHS = 10000
CASH_SEED = 2026
LATE_DAYS = list(range(0, 29))        # the "two weeks late" curve: 0 to 28 days past the reorder point
LATE = 14


def sha(path: str) -> str:
    return hashlib.sha256((ROOT / path).read_bytes()).hexdigest()


def history() -> list[dict]:
    rng = np.random.default_rng(GARLIC["history_seed"])
    cv = GARLIC["history_cv"]
    rows = []
    for m, d0, d1 in HALF_MONTHS:
        days = d1 - d0 + 1
        lam = GARLIC["history_rate"] * math.exp(cv * rng.standard_normal() - cv * cv / 2)
        units = int(rng.poisson(lam * days))
        rows.append({"sku": GARLIC["sku"], "period_start": f"2026-{m:02d}-{d0:02d}", "period_end": f"2026-{m:02d}-{d1:02d}",
                     "days": days, "units_sold": units, "sales": round(units * GARLIC["price"], 2),
                     "avg_sales_price": GARLIC["price"]})
    return rows


def fee_terms(x: dict) -> tuple[float, float]:
    """Amazon's referral rate and fulfilment fee for an example, from the engine's card."""
    f = priors.REFERRAL_BY_CATEGORY[x["category"]]
    fba = priors.fulfilment_fee(x["tier"], x["weight_oz"], x["price"], PRICED_ON, priors.card_named("non_peak"))
    return f, fba


def expected_short(position: int, mu: float, s: float) -> float:
    """E[(D − position)+] for the engine's lead-time demand: the sum of its survival function."""
    total, x = 0.0, position
    while True:
        p = inventory_sim.demand_sf(x, mu, s)
        total += p
        x += 1
        if p < 1e-7:
            return total


def figures() -> dict:
    g = GARLIC
    rows = history()
    f, fba = fee_terms(g)
    contribution = g["price"] * (1 - f) - fba
    margin_unit = contribution - g["landed_cost"]

    # ── the reorder point and the order (inventory_sim.run, exactly as a client's sweep runs it) ──
    data = {
        "sku_economics": rows,
        "inventory_levels": [{"sku": g["sku"], "snapshot_date": "2026-09-30", "fulfillable_quantity": g["on_hand"], "inbound_quantity": 0}],
        "cogs_inputs": [{"sku": g["sku"], "unit_cost": g["landed_cost"], "supplier_lead_time_days": g["lead_days"]}],
        "asin_traffic": [],
    }
    inv = inventory_sim.run(data, np.random.default_rng(0), today=TODAY)[0]
    rate, rate_sd = inv["daily_velocity_mean"], inv["daily_velocity_std"]
    det = inv["details"]
    mu, s = det["lead_demand_log_mu"], det["lead_demand_log_sd"]
    rop, qty = inv["reorder_point"], inv["reorder_qty"]

    # ── the cash path (cashflow.run: the wire schedule, the payouts, 10,000 paths) ──
    sep = rows[-2:]
    sep_units = sum(r["units_sold"] for r in sep)
    margin_rows = [{"sku": g["sku"], "period_start": "2026-09-01", "period_end": "2026-09-30", "units": sep_units,
                    "revenue": round(sep_units * g["price"], 2), "amazon_fees": round(sep_units * (g["price"] * f + fba), 2),
                    "cogs": round(sep_units * g["landed_cost"], 2), "ad_spend_allocated": g["ad_month"]}]
    client = {"cash_on_hand": g["cash"], "monthly_fixed_costs": g["fixed_month"]}
    cash = cashflow.run(client, [inv], margin_rows, np.random.default_rng(CASH_SEED), n_paths=CASH_PATHS, today=TODAY)
    cd = cash["details"]
    wire = cd["wires"][0]
    payout_cycle = cd["payout_cycle_days"]
    reserve_days = channels.payout_reserve_days("amazon") if hasattr(channels, "payout_reserve_days") else 0
    transit_days = channels.payout_transit_days("amazon") if hasattr(channels, "payout_transit_days") else 0
    fee_rate_cash = margin_rows[0]["amazon_fees"] / margin_rows[0]["revenue"]

    # the cash each extra $1,000 of starting balance buys: the paths shift one for one, so the
    # balance that keeps 95% (or the tail mean) above zero is read straight off this run
    need_p5 = g["cash"] - cash["trough_p5"]
    need_es = g["cash"] - cash["trough_expected_shortfall"]
    ruin_at = {}
    for start in (g["cash"] - 1000, g["cash"] - 500, g["cash"] + 500):
        alt = cashflow.run({"cash_on_hand": float(start), "monthly_fixed_costs": g["fixed_month"]}, [inv], margin_rows,
                           np.random.default_rng(CASH_SEED), n_paths=CASH_PATHS, today=TODAY)
        ruin_at[str(int(start - g["cash"]))] = alt["p_ruin"]

    # ── the cash conversion cycle, in the engine's own terms ──
    sell_days = qty / rate                       # the order sells over this many days once it lands
    paid_days = reserve_days + payout_cycle / 2 + transit_days
    ccc = g["lead_days"] + sell_days / 2 + paid_days

    # ── two weeks late (inventory_sim's exact lead-time demand) ──
    late = []
    for d in LATE_DAYS:
        pos = max(0, int(round(rop - d * rate)))
        late.append({"days_late": d, "position": pos, "p_out": round(inventory_sim.demand_sf(pos, mu, s), 4),
                     "units_short": round(expected_short(pos, mu, s), 1)})
    on_time, two_weeks = late[0], late[LATE]
    # what waiting saves: two weeks' interest on the wire (the engine's capital rate) and two weeks'
    # storage on the order, at the rate of the month it would have landed
    from datetime import timedelta
    lands = TODAY + timedelta(days=int(cash["details"]["wires"][0]["day"]) + g["lead_days"]) if cash["details"]["wires"] else TODAY
    carry_saved = (qty * g["landed_cost"] * inventory_econ.ANNUAL_CAPITAL_RATE * LATE / 365
                   + qty * g["volume_cuft"] * fees.storage_rate(lands.month) * LATE / 30)
    # a sale lost is the unit margin lost (the newsvendor's C_u, before the low-inventory fee)
    lost_on_time = on_time["units_short"] * margin_unit
    lost_late = two_weeks["units_short"] * margin_unit

    # ── the service level each SKU earns (inventory_econ.critical_fractile) ──
    cycle = g["lead_days"] + inventory_econ.REVIEW_PERIOD_DAYS
    # each SKU's own row of Amazon's 2026 low-inventory schedule: both are large standard, under 3 lb
    lilf_row = lambda x: fees.low_inventory_fee(0.0, fees.low_inventory_tier("Large standard", x["weight_oz"] / 16)[0])
    cf_g = inventory_econ.critical_fractile(margin_unit, g["landed_cost"], g["volume_cuft"], cycle, TODAY.month,
                                            low_inventory_per_unit=lilf_row(g))
    sp = SPATULA
    f_sp, fba_sp = fee_terms(sp)
    contribution_sp = sp["price"] * (1 - f_sp) - fba_sp
    margin_sp = contribution_sp - sp["landed_cost"]
    cf_s = inventory_econ.critical_fractile(margin_sp, sp["landed_cost"], sp["volume_cuft"], cycle, TODAY.month,
                                            low_inventory_per_unit=lilf_row(sp))
    demand_g = inventory_econ.demand_over_cycle(rate, rate_sd, g["lead_days"], inventory_econ.sku_stream(g["sku"]), 20000)
    demand_s = inventory_econ.demand_over_cycle(sp["rate"], sp["rate_sd"], sp["lead_days"], inventory_econ.sku_stream(sp["sku"]), 20000)
    up_to = lambda d, q: int(math.ceil(float(np.quantile(d, q))))

    # ── hold or liquidate the spatula's excess (inventory_econ.hold_vs_liquidate) ──
    excess = int(max(0.0, sp["position"] - sp["rate"] * inventory_econ.HOLD_HORIZON_DAYS))
    hv = inventory_econ.hold_vs_liquidate(excess, sp["rate"], contribution_sp, sp["price"], sp["volume_cuft"],
                                          sp["age_days"], TODAY)

    r2 = lambda v, d=2: round(float(v), d)
    return {
        "about": "Capital & Cash's worked example: an invented garlic press (The Price Curve's) and an invented spatula set, computed by Hubricon's engine (scripts/learn/capital_cash.py). Invented sales, costs and balance; Amazon's published fees.",
        "today": TODAY.isoformat(),
        "priced_on": PRICED_ON.isoformat(),
        "garlic": {
            "what": g["what"], "price": g["price"], "landed_cost": g["landed_cost"], "referral": f, "fba_fee": r2(fba, 3),
            "contribution": r2(contribution, 4), "margin": r2(margin_unit, 4), "volume_cuft": g["volume_cuft"],
            "lead_days": g["lead_days"], "on_hand": g["on_hand"], "ad_month": g["ad_month"],
            "cash": g["cash"], "fixed_month": g["fixed_month"],
            "history": [{k: r[k] for k in ("period_start", "period_end", "days", "units_sold")} for r in rows],
        },
        "reorder": {
            "rate": rate, "rate_sd": rate_sd, "units_month": r2(rate * 365 / 12, 0),
            "lead_demand_p50": det["demand_percentiles"]["p50"], "lead_demand_p95": det["demand_percentiles"]["p95"],
            "lead_demand_mean": r2(rate * g["lead_days"] * math.exp(0.5 * inventory_sim.LEAD_TIME_CV ** 2), 1),
            "reorder_point": rop, "closed_form_rop": det["closed_form_rop"], "safety_stock": inv["safety_stock"],
            "service_level": inventory_sim.SERVICE_LEVEL, "lead_cv": inventory_sim.LEAD_TIME_CV,
            "reorder_qty": qty, "cover_extra_days": inventory_sim.TARGET_COVER_EXTRA_DAYS,
            "days_to_reorder": wire["day"], "wire": wire["amount"], "cycle_days": max(1, int(qty / rate)),
            "position": g["on_hand"], "days_of_cover": inv["days_of_cover"],
        },
        "cycle": {"lead_days": g["lead_days"], "sell_days": r2(sell_days, 1), "payout_cycle": payout_cycle,
                  "reserve_days": reserve_days, "transit_days": transit_days, "paid_days": r2(paid_days, 1),
                  "payout_note": channels.payout_note("amazon"), "ccc": r2(ccc, 1)},
        "cash": {
            "paths": CASH_PATHS, "horizon_days": cash["horizon_days"], "start": g["cash"], "fixed_month": g["fixed_month"],
            "ad_month": g["ad_month"], "fee_rate": r2(fee_rate_cash, 4), "payout_days": cd["payout_days"],
            "wire_day": wire["day"] + 1, "wire": wire["amount"],
            "trough_day": cash["min_p5_day"], "trough_median": cash["trough_median"], "trough_p5": cash["trough_p5"],
            "trough_es": cash["trough_expected_shortfall"], "trough_es_se": cash["trough_expected_shortfall_se"],
            "p_ruin": cash["p_ruin"], "need_p5": r2(need_p5), "need_es": r2(need_es), "ruin_at": ruin_at,
            "p5": cd["p5"], "p50": cd["p50"], "p95": cd["p95"],
        },
        "late": {"days": LATE, "curve": late, "on_time": on_time, "two_weeks": two_weeks,
                 "lost_on_time": r2(lost_on_time), "lost_late": r2(lost_late), "margin": r2(margin_unit, 4),
                 "carry_saved": r2(carry_saved), "lands": lands.isoformat()},
        "service": {
            "cycle_days": cycle, "capital_rate": inventory_econ.ANNUAL_CAPITAL_RATE,
            "obsolescence_rate": inventory_econ.OBSOLESCENCE_RATE,
            # the 2026 schedule's standard-size rows: the lowest band of the smallest to the highest of the largest
            "lilf_lo": min(min(v.values()) for k, v in fees.LOW_INVENTORY_FEE_PER_UNIT.items() if "standard" in k),
            "lilf_hi": max(max(v.values()) for k, v in fees.LOW_INVENTORY_FEE_PER_UNIT.items() if "standard" in k),
            "storage_offpeak": fees.STORAGE_PER_CUFT["standard"]["offpeak"], "storage_peak": fees.STORAGE_PER_CUFT["standard"]["peak"],
            "garlic": {"c_u": r2(cf_g["c_u"], 4), "c_o": r2(cf_g["c_o"], 4), "q": r2(cf_g["q"], 4),
                       "c_u_parts": {k: r2(v, 4) for k, v in cf_g["c_u_parts"].items()},
                       "c_o_parts": {k: r2(v, 4) for k, v in cf_g["c_o_parts"].items()},
                       "up_to_q": up_to(demand_g, cf_g["q"]), "up_to_95": up_to(demand_g, 0.95)},
            "spatula": {"what": sp["what"], "price": sp["price"], "landed_cost": sp["landed_cost"], "fba_fee": r2(fba_sp, 3),
                        "referral": f_sp, "margin": r2(margin_sp, 4), "rate": sp["rate"],
                        "c_u": r2(cf_s["c_u"], 4), "c_o": r2(cf_s["c_o"], 4), "q": r2(cf_s["q"], 4),
                        "c_u_parts": {k: r2(v, 4) for k, v in cf_s["c_u_parts"].items()},
                        "c_o_parts": {k: r2(v, 4) for k, v in cf_s["c_o_parts"].items()},
                        "up_to_q": up_to(demand_s, cf_s["q"]), "up_to_95": up_to(demand_s, 0.95)},
        },
        "hold": {"position": sp["position"], "rate": sp["rate"], "age_days": sp["age_days"], "excess": excess,
                 "hold_horizon": inventory_econ.HOLD_HORIZON_DAYS, "recovery": inventory_econ.LIQUIDATION_RECOVERY_OF_PRICE,
                 "hold_npv": r2(hv["hold_npv"]), "liquidate": r2(hv["liquidate_value"]), "months": hv["months_to_clear"],
                 "decision": hv["decision"], "per_unit_hold": r2(hv["per_unit_hold"], 4),
                 "contribution": r2(contribution_sp, 4)},
        "engine": {p: sha(p) for p in ENGINE_SOURCES},
    }


# ═══ The spreadsheet ═════════════════════════════════════════════════════════════════════════
# learn/files/hubricon-capital-and-cash.xlsx: the same arithmetic as formulas, opened on the
# example. Where the engine computes a distribution exactly (lead-time demand is a Poisson count
# on a lognormal rate), the sheet uses the lognormal alone, which a spreadsheet can evaluate; the
# check below holds the two within a stated tolerance, and the lessons print the engine's figure.
import random  # noqa: E402
import tempfile  # noqa: E402

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

FILE = ROOT / "learn" / "files" / "hubricon-capital-and-cash.xlsx"
STAMP = ROOT / "scripts" / "learn" / "capital-cash.stamp.json"
INPUT = PatternFill("solid", fgColor="FFF2CC")       # yellow: the reader's cells
TITLE = Font(name="Arial", bold=True, size=16)
HEAD = Font(name="Arial", bold=True, size=11)
BODY = Font(name="Arial", size=10)
BOLD = Font(name="Arial", bold=True, size=10)
NOTE = Font(name="Arial", size=9, italic=True, color="555555")
USD, USD0, PCT, PCT1, NUM1, NUM2, INT = '"$"#,##0.00', '"$"#,##0', "0%", "0.0%", "0.0", "0.00", "#,##0"
SERVICE_ROWS = 12
LATE_ROWS = 29


def put(ws, ref, value, font=BODY, fmt=None, fill=None, bold=False):
    cell = ws[ref]
    cell.value = value
    cell.font = BOLD if bold else font
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    return cell


def inputs(ws, rows, col="B"):
    for r, label, value, fmt in rows:
        put(ws, f"A{r}", label)
        put(ws, f"{col}{r}", value, fmt=fmt, fill=INPUT)


def outputs(ws, rows, bold=()):
    for r, label, formula, fmt in rows:
        put(ws, f"A{r}", label, bold=r in bold)
        put(ws, f"B{r}", formula, fmt=fmt, bold=r in bold)


def wire_sheet(ws, fig: dict) -> None:
    """Lesson 3: the reorder point, the order, the wire and its day."""
    ws.title = "Your wire"
    g, r = fig["garlic"], fig["reorder"]
    put(ws, "A1", "Your wire: when to reorder, how much, and on which day the money leaves", TITLE)
    put(ws, "A2", "Yellow cells are yours. Units a day and its spread come from your own sales: the average and the standard deviation of units a day across periods.", NOTE)
    inputs(ws, [
        (4, "Units sold a day (average)", r["rate"], NUM2), (5, "Spread of units a day (standard deviation)", r["rate_sd"], NUM2),
        (6, "Lead time: order to sellable, days", g["lead_days"], INT), (7, "How much the lead time varies (share)", r["lead_cv"], PCT),
        (8, "Service level for the reorder point", r["service_level"], PCT), (9, "On hand + inbound today, units", g["on_hand"], INT),
        (10, "Landed cost a unit", g["landed_cost"], USD), (11, "Days the order covers after it lands", r["cover_extra_days"], INT),
        (12, "Today's date", TODAY, "yyyy-mm-dd"),
    ])
    outputs(ws, [
        (14, "Spread of the rate, on logs", "=IF(AND(B4>0,B5>0),SQRT(LN(1+(B5/B4)^2)),0)", "0.0000"),
        (15, "Lead-time demand: log centre", "=LN(B4*B6)-B14^2/2", "0.0000"),
        (16, "Lead-time demand: log spread", "=SQRT(B14^2+B7^2)", "0.0000"),
        (17, "Average units sold over the lead time", "=B4*B6*EXP(0.5*B7^2)", NUM1),
        (18, "Reorder point (units)", "=CEILING(LOGINV(B8,B15,B16),1)", INT),
        (19, "Safety stock (units)", "=B18-ROUND(B17,0)", INT),
        (20, "The textbook reorder point, for comparison", "=B4*B6+1.645*SQRT(B6*(B4+B5^2)+(B4*B7*B6)^2)", NUM1),
        (21, "Order (units)", "=CEILING(B4*(B6+B11),1)", INT),
        (22, "The wire", "=B21*B10", USD),
        (23, "Day the wire leaves (day 1 is tomorrow)", "=MAX(0,INT((B9-B18)/B4))+1", INT),
        (24, "Date of the wire", "=B12+B23", "yyyy-mm-dd"),
        (25, "Then a wire every (days)", "=MAX(1,INT(B21/B4))", INT),
    ], bold=(18, 21, 22, 24))
    put(ws, "A27", "The reorder point is the 95th percentile of units sold over the lead time, with the rate and the lead time both uncertain. Hubricon's engine also adds the day-to-day randomness of each sale and computes it exactly; on most listings the two agree within a few units.", NOTE)
    put(ws, "A28", "The textbook formula treats a slow month as noise that averages out over the lead time. It usually reads low.", NOTE)
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 18


def service_sheet(ws, fig: dict, rows: list[dict] | None = None) -> None:
    """Lesson 4: the service level each SKU's margin pays for, and the order-up-to it implies."""
    ws.title = "Service level"
    sv = fig["service"]
    put(ws, "A1", "Service level: how much stock each SKU's margin pays for", TITLE)
    put(ws, "A2", "One SKU a row; yellow cells are yours. Short = the margin lost plus Amazon's low-inventory-level fee; left over = storage for the cycle, the cash tied up and a small write-off.", NOTE)
    inputs(ws, [(3, "This month (1–12): storage is dearer October to December", TODAY.month, INT)])
    heads = ["SKU", "Price", "Referral rate", "Fulfilment fee", "Landed cost", "Cubic feet a unit", "Units a day", "Spread of units a day",
             "Lead time, days", "Low-inventory fee if short", "Contribution a unit", "Margin a unit", "Cycle, days", "Storage for the cycle",
             "Cash tied up", "Write-off", "One left over costs", "One short costs", "Service level", "Order up to", "Order up to at 95%"]
    for i, h in enumerate(heads):
        put(ws, f"{chr(65 + i)}5", h, BOLD)
    g, sp = fig["garlic"], fig["service"]["spatula"]
    example = [
        {"sku": "Garlic press", "price": g["price"], "referral": g["referral"], "fba": g["fba_fee"], "cost": g["landed_cost"], "vol": g["volume_cuft"],
         "rate": fig["reorder"]["rate"], "sd": fig["reorder"]["rate_sd"], "lead": g["lead_days"], "lilf": sv["garlic"]["c_u_parts"]["low_inventory_fee"]},
        {"sku": "Spatula set", "price": sp["price"], "referral": sp["referral"], "fba": sp["fba_fee"], "cost": sp["landed_cost"], "vol": SPATULA["volume_cuft"],
         "rate": SPATULA["rate"], "sd": SPATULA["rate_sd"], "lead": SPATULA["lead_days"], "lilf": sv["spatula"]["c_u_parts"]["low_inventory_fee"]},
    ]
    rows = rows if rows is not None else example
    for i in range(SERVICE_ROWS):
        r = 6 + i
        x = rows[i] if i < len(rows) else None
        for col, key, fmt in [("A", "sku", None), ("B", "price", USD), ("C", "referral", PCT), ("D", "fba", USD), ("E", "cost", USD),
                              ("F", "vol", "0.000"), ("G", "rate", NUM2), ("H", "sd", NUM2), ("I", "lead", INT), ("J", "lilf", USD)]:
            put(ws, f"{col}{r}", x[key] if x else None, fmt=fmt, fill=INPUT)
        f = lambda body: f'=IF(OR($B{r}="",$E{r}=""),"",{body})'
        put(ws, f"K{r}", f"B{r}*(1-C{r})-D{r}", USD)
        ws[f"K{r}"].value = f(f"B{r}*(1-C{r})-D{r}")
        put(ws, f"L{r}", f("K{r}-E{r}".format(r=r)), fmt=USD)
        put(ws, f"M{r}", f(f"I{r}+{inventory_econ.REVIEW_PERIOD_DAYS}"), fmt=INT)
        put(ws, f"N{r}", f(f"IF(OR($B$3=10,$B$3=11,$B$3=12),{fees.STORAGE_PER_CUFT['standard']['peak']},{fees.STORAGE_PER_CUFT['standard']['offpeak']})*F{r}*M{r}/30"), fmt=USD)
        put(ws, f"O{r}", f(f"E{r}*{inventory_econ.ANNUAL_CAPITAL_RATE}*M{r}/365"), fmt=USD)
        put(ws, f"P{r}", f(f"E{r}*{inventory_econ.OBSOLESCENCE_RATE}"), fmt=USD)
        put(ws, f"Q{r}", f(f"N{r}+O{r}+P{r}"), fmt=USD)
        put(ws, f"R{r}", f(f"MAX(0,L{r})+J{r}"), fmt=USD)
        put(ws, f"S{r}", f(f"MIN({inventory_econ.FRACTILE_CEIL},MAX({inventory_econ.FRACTILE_FLOOR},IF(R{r}+Q{r}>0,R{r}/(R{r}+Q{r}),{inventory_econ.FRACTILE_FLOOR})))"), fmt=PCT1)
        # demand over the lead time plus the review week, on the lognormal: the rate's spread and the lead time's
        sr = f"IF(AND(G{r}>0,H{r}>0),SQRT(LN(1+(H{r}/G{r})^2)),0)"
        sl = f"{inventory_econ.LEAD_TIME_CV}*I{r}/M{r}"
        mu = f"LN(G{r}*M{r})-({sr})^2/2"
        sd = f"SQRT(({sr})^2+({sl})^2)"
        put(ws, f"T{r}", f(f"CEILING(LOGINV(S{r},{mu},{sd}),1)"), fmt=INT)
        put(ws, f"U{r}", f(f"CEILING(LOGINV(0.95,{mu},{sd}),1)"), fmt=INT)
    put(ws, f"A{7 + SERVICE_ROWS}", "Service level = one short ÷ (one short + one left over), kept between 50% and 99.5%. Cash tied up at 12% a year; write-off 2% of landed cost a cycle; cycle = lead time + a week.", NOTE)
    put(ws, f"A{8 + SERVICE_ROWS}", "Amazon's low-inventory-level fee (2026) depends on the size tier and the days of supply: see the course, lesson 4. Use the under-14-days rate for the cost of running short.", NOTE)
    ws.column_dimensions["A"].width = 18
    for i in range(1, len(heads)):
        ws.column_dimensions[chr(65 + i)].width = 14


def hold_sheet(ws, fig: dict, case: dict | None = None) -> None:
    """Lesson 5: hold the excess or liquidate it now, on the cash each choice brings from today."""
    ws.title = "Hold or sell"
    h, sv = fig["hold"], fig["service"]["spatula"]
    case = case or {"position": h["position"], "rate": h["rate"], "price": sv["price"], "contribution": h["contribution"],
                    "vol": SPATULA["volume_cuft"], "age": h["age_days"], "month": TODAY.month}
    put(ws, "A1", "Hold or sell: what units already on the shelf are worth, kept or liquidated", TITLE)
    put(ws, "A2", "The landed cost is spent either way, so it counts on neither side. Holding earns each month's sales less storage and the aged surcharge; liquidating returns a share of the price now.", NOTE)
    inputs(ws, [
        (4, "Units on hand", case["position"], INT), (5, "Units sold a day", case["rate"], NUM2), (6, "Price", case["price"], USD),
        (7, "Cash a unit brings when it sells (price less Amazon's fees)", case["contribution"], USD), (8, "Cubic feet a unit", case["vol"], "0.000"),
        (9, "Age of the units today, days", case["age"], INT), (10, "This month (1–12)", case["month"], INT),
        (11, "Days of cover to keep", inventory_econ.HOLD_HORIZON_DAYS, INT), (12, "Liquidation returns (share of price)", inventory_econ.LIQUIDATION_RECOVERY_OF_PRICE, PCT),
        (13, "Interest on cash a year", inventory_econ.ANNUAL_CAPITAL_RATE, PCT),
    ])
    outputs(ws, [
        (15, "Units of cover to keep", "=B5*B11", NUM1),
        (16, "Excess units", "=INT(MAX(0,B4-B15))", INT),
        (17, "Hold: what the excess brings, sold over time", "=SUM(M23:M46)+M47-SUM(X23:X46)-X47", USD),
        (18, "Liquidate now", "=B16*B6*B12", USD),
        (19, "The better choice", '=IF(B16=0,"No excess",IF(B18>B17,"Liquidate","Hold"))', None),
        (20, "Months until the excess is gone", '=IF(B16=0,0,IFERROR(MATCH(TRUE,INDEX(I23:I46<=B15+0.000000001,0),0),24))', INT),
    ], bold=(17, 18, 19))
    # Amazon's storage by month and aged surcharge by age, as the engine carries them
    put(ws, "Z4", "Aged surcharge (a cubic foot a month)", BOLD)
    put(ws, "Z5", "From day", BOLD)
    put(ws, "AA5", "Rate", BOLD)
    put(ws, "Z6", 0)
    put(ws, "AA6", 0, fmt=USD)
    for i, (lo, _hi, rate) in enumerate(fees.AGED_SURCHARGE_PER_CUFT):
        put(ws, f"Z{7 + i}", lo)
        put(ws, f"AA{7 + i}", rate, fmt=USD)
    last_aged = 6 + len(fees.AGED_SURCHARGE_PER_CUFT)
    put(ws, f"Z{last_aged + 2}", "Past 365 days, at least this a unit a month:", NOTE)
    put(ws, f"AA{last_aged + 2}", fees.AGED_SURCHARGE_MIN_PER_UNIT_365_PLUS, fmt=USD)
    put(ws, f"Z{last_aged + 3}", "Storage a cubic foot: Jan–Sep / Oct–Dec", NOTE)
    put(ws, f"AA{last_aged + 3}", fees.STORAGE_PER_CUFT["standard"]["offpeak"], fmt=USD)
    put(ws, f"AB{last_aged + 3}", fees.STORAGE_PER_CUFT["standard"]["peak"], fmt=USD)
    aged = f"$Z$6:$AA${last_aged}"
    floor_ref, off_ref, peak_ref = f"$AA${last_aged + 2}", f"$AA${last_aged + 3}", f"$AB${last_aged + 3}"
    # the position modelled is the excess plus the cover, as the engine's (hold_vs_liquidate)
    for side, c0, title in (("all", "D", "Holding everything: month by month"), ("cover", "O", "Liquidating the excess, holding the cover")):
        cols = [chr(ord(c0) + k) for k in range(10)]   # month, cal month, age, start, sold, left after, carry, cash, discount, value
        put(ws, f"{cols[0]}21", title, HEAD)
        for c, head in zip(cols, ["Month", "Calendar month", "Age, days", "Units at the start", "Sold", "Left after", "Storage and surcharge", "Cash", "Discount", "Worth today"]):
            put(ws, f"{c}22", head, BOLD)
        for m in range(1, 25):
            r = 22 + m
            mo, cm, age, start, sold, after, carry, cash, disc, val = (f"{c}{r}" for c in cols)
            put(ws, mo, m)
            put(ws, cm, f"=MOD($B$10+{m}-1,12)+1")
            put(ws, age, f"=$B$9+30*{m}")
            first = "$B$16+$B$15" if side == "all" else "$B$15"
            put(ws, start, f"={first}" if m == 1 else f"={cols[5]}{r - 1}", fmt=NUM1)
            put(ws, sold, f"=MIN({start},$B$5*30)", fmt=NUM1)
            put(ws, after, f"=MAX(0,{start}-{sold})", fmt=NUM1)
            surcharge = f"VLOOKUP({age},{aged},2,TRUE)"
            put(ws, carry, f"=IF({age}>365,MAX({start}*$B$8*(IF({cm}>=10,{peak_ref},{off_ref})+{surcharge}),{start}*{floor_ref}),{start}*$B$8*(IF({cm}>=10,{peak_ref},{off_ref})+{surcharge}))", fmt=USD)
            put(ws, cash, f"={sold}*$B$7-{carry}", fmt=USD)
            put(ws, disc, f"=(1+$B$13/12)^{m}", fmt="0.0000")
            put(ws, val, f"={cash}/{disc}", fmt=USD)
        # what is left after 24 months is liquidated then
        put(ws, f"{cols[0]}47", "Left at 24 months, liquidated then")
        put(ws, f"{cols[9]}47", f"={cols[5]}46*$B$6*$B$12/(1+$B$13/12)^24", fmt=USD)
    put(ws, "A49", "Hold = (everything held) − (the cover alone held): what the excess adds by staying. Discounted at the interest rate, a month at a time. Leftovers after 24 months are liquidated.", NOTE)
    ws.column_dimensions["A"].width = 56
    ws.column_dimensions["B"].width = 16


def late_sheet(ws, fig: dict) -> None:
    """Lesson 7: the chance of running out, and the units short, by days late."""
    ws.title = "Two weeks late"
    put(ws, "A1", "Two weeks late: what waiting past the reorder point costs", TITLE)
    put(ws, "A2", "Linked to Your wire. The chance of running out before the order lands and the units you would be short, for each day the wire goes late.", NOTE)
    inputs(ws, [(4, "Margin a unit (price less fees and landed cost)", fig["garlic"]["margin"], USD)])
    put(ws, "A5", "From Your wire: units a day, reorder point, and the lead-time demand's log centre and spread")
    for c, ref in zip("BCDE", ["'Your wire'!B4", "'Your wire'!B18", "'Your wire'!B15", "'Your wire'!B16"]):
        put(ws, f"{c}5", f"={ref}", fmt="0.0000")
    for c, head in zip("ABCDE", ["Days late", "Position when the wire goes", "Chance of running out", "Units short, on average", "Margin lost a cycle"]):
        put(ws, f"{c}7", head, BOLD)
    for d in range(LATE_ROWS):
        r = 8 + d
        put(ws, f"A{r}", d)
        put(ws, f"B{r}", f"=MAX(1,ROUND($C$5-A{r}*$B$5,0))", fmt=INT)
        put(ws, f"C{r}", f"=1-LOGNORMDIST(B{r},$D$5,$E$5)", fmt=PCT1)
        put(ws, f"D{r}", f"=EXP($D$5+$E$5^2/2)*NORMSDIST(($D$5+$E$5^2-LN(B{r}))/$E$5)-B{r}*NORMSDIST(($D$5-LN(B{r}))/$E$5)", fmt=NUM1)
        put(ws, f"E{r}", f"=D{r}*$B$4", fmt=USD)
    put(ws, f"A{9 + LATE_ROWS}", "Left out: the sales rank a stockout loses, and Amazon's low-inventory-level fee while cover is thin. Both make late worse.", NOTE)
    for c, w in zip("ABCDE", [44, 26, 22, 22, 20]):
        ws.column_dimensions[c].width = w


CASH_PATHS_SHEET = 1000


def cycle_sheet(ws, fig: dict) -> None:
    """Lesson 1: the days a dollar is gone, and the cash the cycle holds."""
    ws.title = "Your cycle"
    cy, g = fig["cycle"], fig["garlic"]
    put(ws, "A1", "Your cycle: how many days a dollar wired to your supplier is gone", TITLE)
    put(ws, "A2", "Yellow cells are yours. Amazon holds a sale's money until 7 days after delivery (DD+7), settles every 14 days, and the bank takes a few days more.", NOTE)
    inputs(ws, [
        (4, "Lead time: order to sellable, days", "='Your wire'!B6", INT), (5, "Units sold a day", "='Your wire'!B4", NUM2),
        (6, "Units in the order", "='Your wire'!B21", INT), (7, "Landed cost a unit", "='Your wire'!B10", USD),
        (8, "Days from a sale until the platform can pay it", cy["reserve_days"], INT), (9, "Days between payouts", cy["payout_cycle"], INT),
        (10, "Days for the transfer to reach your bank", cy["transit_days"], INT),
        (11, "Share of the order paid on the day you order", 1, PCT), (12, "Days after the order the rest is paid", 0, INT),
    ])
    outputs(ws, [
        (14, "Days the order takes to sell", "=B6/B5", NUM1),
        (15, "Days until the average unit sells", "=B4+B14/2", NUM1),
        (16, "Days until you are paid for it, on average", "=B8+B9/2+B10", NUM1),
        (17, "Days of supplier credit", "=(1-B11)*B12", NUM1),
        (18, "Days a dollar is gone", "=B15+B16-B17", NUM1),
        (19, "Cash your cycle holds", "=B5*B7*B18", USD0),
    ], bold=(18, 19))
    put(ws, "A21", "Amazon: a sale is payable 7 days after delivery (its own example: delivered January 6, available January 14), settles every 14 days unless you request a payout, and the transfer takes up to five business days. Shopify Payments: 3 to 5 business days, on a daily, weekly or monthly schedule.", NOTE)
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 16


def cash_sheet(ws, fig: dict) -> None:
    """Lessons 2 and 6: the bank balance day by day, its low point, and the bad case."""
    ws.title = "Cash path"
    c, cy = fig["cash"], fig["cycle"]
    put(ws, "A1", "Cash path: your bank balance for the next 90 days, and how much cash to hold", TITLE)
    put(ws, "A2", "Yellow cells are yours. The wire comes from Your wire. The bad case is simulated: press F9 (Excel) or Ctrl+Shift+F9 (LibreOffice) for new draws; the figures barely move.", NOTE)
    inputs(ws, [
        (4, "Cash in the bank today", c["start"], USD0), (5, "Fixed costs a month", c["fixed_month"], USD0), (6, "Ad spend a month", c["ad_month"], USD0),
        (7, "Units sold a day", "='Your wire'!B4", NUM2), (8, "Spread of units a day", "='Your wire'!B5", NUM2),
        (9, "Price", fig["garlic"]["price"], USD), (10, "Amazon's fees as a share of revenue", c["fee_rate"], PCT1),
        (11, "Day the wire leaves", "='Your wire'!B23", INT), (12, "The wire", "='Your wire'!B22", USD0),
        (13, "Days between payouts", cy["payout_cycle"], INT),
    ])
    outputs(ws, [
        (15, "Cash a day of sales brings (after fees)", "=B7*B9*(1-B10)", USD),
        (16, "Each payout, on average", "=B13*B15", USD0),
        (17, "Low point of the balance (middle case)", "=MIN(I6:I95)", USD0),
        (18, "Day of the low point", "=MATCH(B17,I6:I95,0)", INT),
        (19, "Low point, 1 path in 20 is worse than", f"=PERCENTILE(AA6:AA{5 + CASH_PATHS_SHEET},0.05)", USD0),
        (20, "Average of the worst 20th", f'=AVERAGEIF(AA6:AA{5 + CASH_PATHS_SHEET},"<="&B19)', USD0),
        (21, "Cash to hold for this cycle (before your floor)", "=B4-B20", USD0),
    ], bold=(17, 19, 20, 21))
    put(ws, "A23", "Each payout carries a payout cycle's sales, made two weeks or so before it lands: the sales the platform holds today pay the first ones. Fixed costs and ads leave every day; the wire leaves on its day.", NOTE)
    for col, head in zip("DEFGHI", ["Day", "Payout in", "Fixed costs", "Ads", "Wire", "Balance"]):
        put(ws, f"{col}5", head, BOLD)
    for d in range(1, 91):
        r = 5 + d
        put(ws, f"D{r}", d)
        put(ws, f"E{r}", f"=IF(MOD(D{r},$B$13)=0,$B$16,0)", fmt=USD0)
        put(ws, f"F{r}", "=$B$5/30", fmt=USD0)
        put(ws, f"G{r}", "=$B$6/30", fmt=USD0)
        put(ws, f"H{r}", f"=IF(D{r}=$B$11,$B$12,0)", fmt=USD0)
        prev = "$B$4" if d == 1 else f"I{r - 1}"
        put(ws, f"I{r}", f"={prev}+E{r}-F{r}-G{r}-H{r}", fmt=USD0)
    # the bad case: each payout's sales drawn around their average (units a day vary day to day,
    # and each day's count is random), the balance read on the eve of every payout and on day 90
    landings = [k for k in range(14, 91, 14)]
    put(ws, "K4", "The bad case, simulated: one path a row", HEAD)
    put(ws, "K5", "Path", BOLD)
    pay_cols = [chr(ord("L") + j) for j in range(len(landings))]           # payouts
    eve_cols = [chr(ord("L") + len(landings) + j) for j in range(len(landings) + 1)]   # balances
    for j, col in enumerate(pay_cols):
        put(ws, f"{col}5", f"Payout {j + 1}", BOLD)
    for j, col in enumerate(eve_cols):
        put(ws, f"{col}5", f"Before payout {j + 1}" if j < len(landings) else "Day 90", BOLD)
    put(ws, "AA5", "Lowest", BOLD)
    sd = "SQRT($B$13*($B$7+$B$8^2))*$B$9*(1-$B$10)"
    for p_ in range(CASH_PATHS_SHEET):
        r = 6 + p_
        put(ws, f"K{r}", p_ + 1)
        for j, col in enumerate(pay_cols):
            put(ws, f"{col}{r}", f"=$B$16+NORMINV(RAND(),0,{sd})", fmt=USD0)
        for j, col in enumerate(eve_cols):
            day = landings[j] - 1 if j < len(landings) else 90
            paid = "+".join(f"{pay_cols[i]}{r}" for i in range(len(landings)) if landings[i] <= day) or "0"
            put(ws, f"{col}{r}", f"=$B$4-{day}*($B$5+$B$6)/30-IF($B$11<={day},$B$12,0)+{paid}", fmt=USD0)
        put(ws, f"AA{r}", f"=MIN({eve_cols[0]}{r}:{eve_cols[-1]}{r})", fmt=USD0)
    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["B"].width = 16


def start_sheet(ws, fig: dict) -> None:
    ws.title = "Start here"
    lines = [
        ("Capital & Cash", TITLE),
        ("A free course from Hubricon: hubricon.com/learn/capital-and-cash", BODY),
        ("", BODY),
        ("What this file does", HEAD),
        ("Your cycle: how many days a dollar wired to your supplier is gone, and the cash that ties up.", BODY),
        ("Cash path: your bank balance for 90 days, its low point after the wire, and the cash to hold for a bad month.", BODY),
        ("Your wire: the reorder point, the order, and the day and size of the wire, from your own sales and lead time.", BODY),
        ("Service level: how much stock each SKU's margin pays for, one SKU a row.", BODY),
        ("Hold or sell: what units already on the shelf bring, kept or liquidated.", BODY),
        ("Two weeks late: the chance of running out, and the margin lost, for each day a wire goes late.", BODY),
        ("", BODY),
        ("How to use it", HEAD),
        ("Yellow cells are yours. Everything else is a formula you can read.", BODY),
        (f"It opens on the course's examples: {fig['garlic']['what']} and {SPATULA['what']}. Their sales, costs and bank balance are invented; the fees are Amazon's published 2026 schedule. Type over them.", BODY),
        ("", BODY),
        ("Where the arithmetic comes from", HEAD),
        ("Hubricon's engine: models/inventory_sim.py, inventory_econ.py and cashflow.py. This file was checked against it before it was published.", BODY),
        ("One difference, stated: the engine counts lead-time demand exactly, sale by sale; a spreadsheet uses the lognormal that carries it, which agrees within a few units on most listings.", BODY),
    ]
    for i, (text, font) in enumerate(lines):
        put(ws, f"A{i + 1}", text, font)
    ws.column_dimensions["A"].width = 120


def build(path: Path, fig: dict, service_rows: list[dict] | None = None) -> Workbook:
    wb = Workbook()
    start_sheet(wb.active, fig)
    cycle_sheet(wb.create_sheet(), fig)
    cash_sheet(wb.create_sheet(), fig)
    wire_sheet(wb.create_sheet(), fig)
    service_sheet(wb.create_sheet(), fig, service_rows)
    hold_sheet(wb.create_sheet(), fig)
    late_sheet(wb.create_sheet(), fig)
    wb.properties.creator = "Hubricon"
    wb.properties.title = "Capital & Cash"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return wb


def recalc(path: Path) -> None:
    sys.path.insert(0, str(ROOT / "scripts" / "learn"))
    from verify_fee_staircase import recalc as lo_recalc
    lo_recalc(path)


# ═══ The check: the sheet against the engine ═══════════════════════════════════════════════════
def _wire_case(rng: random.Random) -> dict:
    rate = rng.uniform(1.0, 60.0)
    return {"rate": rate, "sd": rate * rng.uniform(0.08, 0.45), "lead": rng.choice([30, 45, 60, 75, 90, 120]),
            "position": int(rate * rng.uniform(20, 200)), "cost": round(rng.uniform(1.5, 40.0), 2)}


def _service_case(rng: random.Random, i: int) -> dict:
    price = round(rng.uniform(9.0, 80.0), 2)
    cost = round(price * rng.uniform(0.12, 0.5), 2)
    fba = round(rng.uniform(3.0, 9.0), 3)
    rate = rng.uniform(0.5, 40.0)
    return {"sku": f"Case {i + 1}", "price": price, "referral": rng.choice([0.08, 0.15, 0.17]), "fba": fba, "cost": cost,
            "vol": round(rng.uniform(0.01, 0.6), 3), "rate": rate, "sd": rate * rng.uniform(0.1, 0.5),
            "lead": rng.choice([30, 45, 60, 75, 90]), "lilf": None}


def _hold_case(rng: random.Random) -> dict:
    rate = rng.uniform(0.3, 12.0)
    price = round(rng.uniform(8.0, 70.0), 2)
    return {"position": int(rate * rng.uniform(100, 700)), "rate": rate, "price": price,
            "contribution": round(price * rng.uniform(0.25, 0.7), 4), "vol": round(rng.uniform(0.01, 1.2), 3),
            "age": rng.choice([60, 120, 190, 230, 280, 330, 380]), "month": rng.randint(1, 12)}


def verify(fig: dict, cases: int = 24) -> int:
    rng = random.Random(20261002)
    problems: list[str] = []
    wires = [_wire_case(rng) for _ in range(cases)]
    service = [_service_case(rng, i) for i in range(SERVICE_ROWS)]
    holds = [_hold_case(rng) for _ in range(cases)]
    for row in service:   # the low-inventory fee the engine charges for being short, so the rows compare like for like
        row["lilf"] = inventory_econ.critical_fractile(1.0, row["cost"], row["vol"], row["lead"] + 7, TODAY.month)["c_u_parts"]["low_inventory_fee"]
    with tempfile.TemporaryDirectory(prefix="hubricon-cc-") as tmp:
        path = Path(tmp) / "check.xlsx"
        wb = build(path, fig, service_rows=service)
        for k, c in enumerate(wires):
            ws = wb.copy_worksheet(wb["Your wire"])
            ws.title = f"W{k}"
            for ref, v in (("B4", c["rate"]), ("B5", c["sd"]), ("B6", c["lead"]), ("B9", c["position"]), ("B10", c["cost"])):
                ws[ref].value = v
        for k, c in enumerate(holds):
            ws = wb.copy_worksheet(wb["Hold or sell"])
            ws.title = f"H{k}"
            for ref, key in (("B4", "position"), ("B5", "rate"), ("B6", "price"), ("B7", "contribution"), ("B8", "vol"), ("B9", "age"), ("B10", "month")):
                ws[ref].value = c[key]
        wb.save(path)
        recalc(path)
        got = load_workbook(path, data_only=True)

        for k, c in enumerate(wires):
            ws = got[f"W{k}"]
            mu, s = inventory_sim.lead_time_mixture(c["rate"], c["sd"], c["lead"])
            rop = inventory_sim.demand_quantile(inventory_sim.SERVICE_LEVEL, mu, s)
            exact = [("closed-form ROP", ws["B20"].value, inventory_sim.closed_form_rop(c["rate"], c["sd"], c["lead"]), 1e-6),
                     ("order", ws["B21"].value, math.ceil(c["rate"] * (c["lead"] + inventory_sim.TARGET_COVER_EXTRA_DAYS)), 0),
                     ("log centre", ws["B15"].value, mu, 1e-9), ("log spread", ws["B16"].value, s, 1e-9)]
            for name, a, b, tol in exact:
                if abs(float(a) - float(b)) > tol * max(1.0, abs(float(b))):
                    problems.append(f"W{k}: {name} sheet {a} engine {b}")
            # the reorder point: lognormal against the engine's exact Poisson–lognormal count
            if abs(ws["B18"].value - rop) > max(3, 0.015 * rop):
                problems.append(f"W{k}: reorder point sheet {ws['B18'].value} engine {rop}")
            wire_day = max(0, int((c["position"] - ws["B18"].value) / c["rate"])) + 1
            if ws["B23"].value != wire_day:
                problems.append(f"W{k}: days to the wire sheet {ws['B23'].value} expected {wire_day}")

        sv = got["Service level"]
        for i, c in enumerate(service):
            r = 6 + i
            contribution = c["price"] * (1 - c["referral"]) - c["fba"]
            cf = inventory_econ.critical_fractile(contribution - c["cost"], c["cost"], c["vol"], c["lead"] + 7, TODAY.month)
            for name, col, want in (("one short", "R", cf["c_u"]), ("one left over", "Q", cf["c_o"]), ("service level", "S", cf["q"])):
                if abs(float(sv[f"{col}{r}"].value) - want) > 1e-9 * max(1.0, abs(want)):
                    problems.append(f"Service row {i + 1}: {name} sheet {sv[f'{col}{r}'].value} engine {want}")
            demand = inventory_econ.demand_over_cycle(c["rate"], c["sd"], c["lead"], inventory_econ.sku_stream(c["sku"]), 20000)
            up = math.ceil(float(np.quantile(demand, cf["q"])))
            if abs(sv[f"T{r}"].value - up) > max(3, 0.03 * up):
                problems.append(f"Service row {i + 1}: order up to sheet {sv[f'T{r}'].value} engine {up}")

        for k, c in enumerate(holds):
            ws = got[f"H{k}"]
            excess = int(max(0.0, c["position"] - c["rate"] * inventory_econ.HOLD_HORIZON_DAYS))
            today = date(2026, c["month"], 1)
            if excess == 0:
                if ws["B16"].value != 0:
                    problems.append(f"H{k}: excess sheet {ws['B16'].value} engine 0")
                continue
            hv = inventory_econ.hold_vs_liquidate(excess, c["rate"], c["contribution"], c["price"], c["vol"], c["age"], today)
            for name, ref, want in (("excess", "B16", excess), ("hold", "B17", hv["hold_npv"]), ("liquidate", "B18", hv["liquidate_value"])):
                if abs(float(ws[ref].value) - want) > 1e-6 * max(1.0, abs(want)):
                    problems.append(f"H{k}: {name} sheet {ws[ref].value} engine {want}")
            if (ws["B19"].value or "").lower() != hv["decision"]:
                problems.append(f"H{k}: decision sheet {ws['B19'].value} engine {hv['decision']}")
            if ws["B20"].value != hv["months_to_clear"]:
                problems.append(f"H{k}: months sheet {ws['B20'].value} engine {hv['months_to_clear']}")

        late = got["Two weeks late"]
        for p in fig["late"]["curve"]:
            r = 8 + p["days_late"]
            if abs(late[f"C{r}"].value - p["p_out"]) > 0.012:
                problems.append(f"Late {p['days_late']}d: chance sheet {late[f'C{r}'].value:.4f} engine {p['p_out']}")
            if abs(late[f"D{r}"].value - p["units_short"]) > max(1.5, 0.1 * p["units_short"]):
                problems.append(f"Late {p['days_late']}d: units short sheet {late[f'D{r}'].value:.1f} engine {p['units_short']}")

        cash, c = got["Cash path"], fig["cash"]
        if abs(cash["B17"].value - c["trough_median"]) > 0.02 * c["start"]:
            problems.append(f"Cash: middle-case low point sheet {cash['B17'].value:.0f} engine {c['trough_median']}")
        if cash["B18"].value != c["trough_day"]:
            problems.append(f"Cash: day of the low point sheet {cash['B18'].value} engine {c['trough_day']}")
        for name, ref, want in (("bad case (5th percentile)", "B19", c["trough_p5"]), ("average of the worst 20th", "B20", c["trough_es"])):
            if abs(cash[ref].value - want) > 0.03 * c["start"]:
                problems.append(f"Cash: {name} sheet {cash[ref].value:.0f} engine {want}")
        cyc, cy = got["Your cycle"], fig["cycle"]
        if abs(cyc["B18"].value - cy["ccc"]) > 0.051:
            problems.append(f"Cycle: days a dollar is gone sheet {cyc['B18'].value} figures {cy['ccc']}")
        wire = got["Your wire"]
        r = fig["reorder"]
        for name, a, b in (("example order", wire["B21"].value, r["reorder_qty"]), ("example wire", round(wire["B22"].value, 2), r["wire"])):
            if abs(float(a) - float(b)) > 0.005:
                problems.append(f"Example: {name} sheet {a} engine {b}")
        if abs(wire["B18"].value - r["reorder_point"]) > 3:
            problems.append(f"Example: reorder point sheet {wire['B18'].value} engine {r['reorder_point']}")

    for p in problems:
        print("  ✗", p)
    n = 2 * cases + SERVICE_ROWS + len(fig["late"]["curve"])
    print(f"{n} golden cases and the worked example: {'all agree' if not problems else f'{len(problems)} disagree'}")
    return len(problems)


if __name__ == "__main__":
    out = figures()
    FIGURES.write_text(json.dumps(out, indent=1) + "\n")
    c, r, sv, h, l = out["cash"], out["reorder"], out["service"], out["hold"], out["late"]
    print(f"rate {r['rate']}/day (sd {r['rate_sd']}), ROP {r['reorder_point']} (closed form {r['closed_form_rop']}), qty {r['reorder_qty']}, wire ${r['wire']:,} on day {c['wire_day']}")
    print(f"cycle {out['cycle']}")
    print(f"cash: trough day {c['trough_day']}, median {c['trough_median']}, p5 {c['trough_p5']}, ES {c['trough_es']}, ruin {c['p_ruin']}, need {c['need_p5']}/{c['need_es']}, ruin_at {c['ruin_at']}")
    print(f"late: on time {l['on_time']}, two weeks {l['two_weeks']}, lost {l['lost_on_time']} -> {l['lost_late']}")
    print(f"service: garlic q {sv['garlic']['q']} ({sv['garlic']['up_to_q']} vs {sv['garlic']['up_to_95']}), spatula q {sv['spatula']['q']} ({sv['spatula']['up_to_q']} vs {sv['spatula']['up_to_95']}), spatula margin {sv['spatula']['margin']}")
    print(f"hold: excess {h['excess']}, hold {h['hold_npv']} vs liquidate {h['liquidate']}, {h['months']} months, {h['decision']}")
    print(f"wrote {FIGURES.relative_to(ROOT)}")
    if "--figures-only" in sys.argv:
        raise SystemExit(0)
    if verify(out):
        raise SystemExit(1)
    if "--publish" in sys.argv:
        build(FILE, out)
        recalc(FILE)
        STAMP.write_text(json.dumps({
            "file": str(FILE.relative_to(ROOT)),
            "file_sha256": hashlib.sha256(FILE.read_bytes()).hexdigest(),
            "engine_sha256": out["engine"],
            "verified_on": date.today().isoformat(),
        }, indent=2) + "\n")
        print(f"published {FILE.relative_to(ROOT)} and {STAMP.relative_to(ROOT)}")
